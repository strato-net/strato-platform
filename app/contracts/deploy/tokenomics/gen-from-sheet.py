#!/usr/bin/env python3
"""Derive the on-chain tokenomics schedule from the tokenomics spreadsheet.

The spreadsheet is the source of truth. This reads it and writes:

  schedule.json                                 every category's total (wei) and schedule
                                                parameters (TGE and addresses live in
                                                deployment.json, not here)
  ../../tests/Tokenomics/TokenomicsSheet.test.sol
                                                checks every category's contract against the
                                                sheet's month-by-month unlock table

Usage (needs openpyxl):
  python3 gen-from-sheet.py path/to/tokenomics.xlsx

Amounts: each category is the sheet's percentage of Total Network Tokens rounded to
whole tokens; the ecosystem fund is the balancing row (as in the sheet), so the
categories sum to exactly the total supply. Emission rows take each year's amount
from the sheet's monthly values; the last year absorbs rounding.
"""
import argparse
import json
import os
from decimal import Decimal, ROUND_HALF_EVEN, getcontext

import openpyxl

getcontext().prec = 60
WEI = Decimal(10) ** 18
HERE = os.path.dirname(os.path.abspath(__file__))
TEST_OUT = os.path.join(HERE, "..", "..", "tests", "Tokenomics", "TokenomicsSheet.test.sol")
MONTHS = 121           # months 0..120 (10 years after TGE)
TOLERANCE_WEI = 10**12  # 0.000001 STRATO: float noise in the sheet, not schedule error

# Settings row -> (key, Unlocking Schedule row). Order matches the sheet.
ROWS = [
    (14, "team", 40),
    (15, "advisors", 41),
    (16, "investors", 42),
    (17, "publicSales", 43),
    (18, "airdrop", 44),
    (19, "tvlDeal1", 45),
    (20, "ongoingTvl", 46),
    (21, "bootstrapping", 47),
    (22, "ecosystemFund", 48),
    (23, "ecosystemDev", 49),
]
BALANCING = "ecosystemFund"


def dec(x):
    return Decimal(repr(x))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    settings = wb["Tokenomics Settings"]
    sched = wb["Token Unlocking Schedule"]
    assert sched["B5"].value == "OK", "the sheet's own checks do not pass"

    supply_tokens = dec(settings["C9"].value)
    first = 6  # column G = month 0
    for m in range(MONTHS):
        assert sched.cell(row=16, column=first + 1 + m).value == m, "month row out of place"

    cats = []
    for srow, key, urow in ROWS:
        label = settings[f"B{srow}"].value
        assert sched[f"C{urow}"].value == label, f"row mismatch for {label}"
        pct = dec(settings[f"D{srow}"].value)
        monthly = [dec(sched.cell(row=urow, column=first + 1 + m).value or 0) for m in range(MONTHS)]
        sheet_total = dec(sched[f"E{urow}"].value)
        cat = {
            "key": key,
            "label": label,
            "category": settings[f"C{srow}"].value,
            "sheetPercent": str(pct * 100),
            "monthly": monthly,
            "sheetTotal": sheet_total,
        }
        if settings[f"E{srow}"].value == "Modelled":
            cat["kind"] = "emission"
        else:
            assert settings[f"H{srow}"].value in (0, 1), f"{label}: only monthly periodicity is supported"
            bps = dec(settings[f"F{srow}"].value) * 10000
            assert bps == bps.to_integral_value(), f"{label}: cliff unlock is not whole bps"
            cat.update(kind="vesting", cliffMonths=int(settings[f"E{srow}"].value),
                       cliffUnlockBps=int(bps), vestingMonths=int(settings[f"G{srow}"].value))
        if key != BALANCING:
            cat["tokens"] = (pct * supply_tokens).quantize(Decimal(1), rounding=ROUND_HALF_EVEN)
        cats.append(cat)

    fund = next(c for c in cats if c["key"] == BALANCING)
    fund["tokens"] = supply_tokens - sum(c["tokens"] for c in cats if c is not fund)
    supply_wei = int(supply_tokens * WEI)
    for c in cats:
        c["wei"] = int(c["tokens"] * WEI)
        if c["kind"] == "emission":
            assert c["monthly"][0] == 0, f"{c['label']}: emits at TGE"
            years = [sum(c["monthly"][12 * y + 1:12 * y + 13]) for y in range(10)]
            yearly = [int((y * WEI).quantize(Decimal(1), rounding=ROUND_HALF_EVEN)) for y in years]
            yearly[-1] = c["wei"] - sum(yearly[:-1])
            c["yearlyWei"] = yearly
    assert sum(c["wei"] for c in cats) == supply_wei

    # Expected cumulative unlock per month: the sheet's own table, scaled to the
    # rounded category total.
    for c in cats:
        cum, out = Decimal(0), []
        for v in c["monthly"]:
            cum += v
            out.append(int((cum / c["sheetTotal"] * c["wei"]).quantize(Decimal(1), rounding=ROUND_HALF_EVEN)))
        c["expected"] = out

    write_json(cats, supply_wei, args, settings)
    write_test(cats, supply_wei)


def write_json(cats, supply_wei, args, settings):
    out = {
        "source": os.path.basename(args.xlsx),
        "sheetTgeDate": settings["C8"].value.strftime("%Y-%m-%d"),
        "totalSupplyWei": str(supply_wei),
        "categories": [],
    }
    for c in cats:
        e = {"key": c["key"], "label": c["label"], "category": c["category"], "kind": c["kind"],
             "sheetPercent": c["sheetPercent"], "tokens": str(c["tokens"]), "wei": str(c["wei"])}
        if c["kind"] == "vesting":
            e.update(cliffMonths=c["cliffMonths"], cliffUnlockBps=c["cliffUnlockBps"], vestingMonths=c["vestingMonths"])
        else:
            e["yearlyWei"] = [str(y) for y in c["yearlyWei"]]
        out["categories"].append(e)
    with open(os.path.join(HERE, "schedule.json"), "w") as f:
        json.dump(out, f, indent=2)
        f.write("\n")


def write_test(cats, supply_wei):
    L = []
    w = L.append
    w("// GENERATED by app/contracts/deploy/tokenomics/gen-from-sheet.py from the tokenomics")
    w("// spreadsheet. Do not edit; regenerate.")
    w("//")
    w("// Each category's contract must reproduce the sheet's cumulative unlock table for")
    w(f"// months 0..{MONTHS - 1} to within {TOLERANCE_WEI} wei (float noise in the sheet).")
    w('import "../../concrete/Tokenomics/MonthlyVestingWallet.sol";')
    w('import "../../concrete/Tokenomics/EmissionVestingWallet.sol";')
    w("")
    w("contract Describe_TokenomicsSheet {")
    w("    address constant BENEFICIARY = address(0x7e5);")
    w("    uint64 constant TGE = 1792972800;")
    w(f"    uint constant TOLERANCE = {TOLERANCE_WEI};")
    w("")
    w("    function _check(string name, uint month, uint got, uint want) internal {")
    w("        uint diff = got > want ? got - want : want - got;")
    w('        require(diff <= TOLERANCE, name + " month " + string(month) + ": got " + string(got) + ", sheet " + string(want));')
    w("    }")
    w("")
    for c in cats:
        name = c["key"][0].upper() + c["key"][1:]
        w("")
        w(f"    // {c['label']}: {c['tokens']} STRATO")
        w(f"    function _expected{name}() internal returns (uint[]) {{")
        w(f"        uint[] memory e = new uint[]({MONTHS});")
        for m, v in enumerate(c["expected"]):
            if v:
                w(f"        e[{m}] = {v};")
        w("        return e;")
        w("    }")
        w("")
        w(f"    function it_matches_the_sheet_for_{c['key']}() public {{")
        w(f"        uint[] memory want = _expected{name}();")
        if c["kind"] == "vesting":
            w(f'        MonthlyVestingWallet s = new MonthlyVestingWallet("{c["label"]}", BENEFICIARY, TGE, {c["cliffMonths"]}, {c["cliffUnlockBps"]}, {c["vestingMonths"]});')
            w(f"        for (uint m = 0; m < {MONTHS}; m++) {{ _check(\"{c['key']}\", m, s.cumulativeUnlocked({c['wei']}, m), want[m]); }}")
        else:
            w(f"        uint[] memory yearly = new uint[]({len(c['yearlyWei'])});")
            for i, y in enumerate(c["yearlyWei"]):
                w(f"        yearly[{i}] = {y};")
            w(f'        EmissionVestingWallet s = new EmissionVestingWallet("{c["label"]}", BENEFICIARY, TGE, yearly);')
            w(f'        require(s.total() == {c["wei"]}, "emission total");')
            w(f"        for (uint m = 0; m < {MONTHS}; m++) {{ _check(\"{c['key']}\", m, s.cumulativeUnlocked(m), want[m]); }}")
        w(f'        _check("{c["key"]} fully unlocked", {MONTHS - 1}, {c["wei"]}, want[{MONTHS - 1}]);')
        w("    }")
    w("")
    w("    function it_sums_to_the_total_supply() public {")
    w(f'        require({" + ".join(str(c["wei"]) for c in cats)} == {supply_wei}, "categories sum to the total supply");')
    w("    }")
    w("}")
    with open(TEST_OUT, "w") as f:
        f.write("\n".join(L) + "\n")


if __name__ == "__main__":
    main()
